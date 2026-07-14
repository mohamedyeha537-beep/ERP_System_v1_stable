import openpyxl
import os
import zipfile
import re

path = r"c:\Users\PIXEL\OneDrive\Desktop\سيف\New folder\Bill of Material (mrp.bom) (1).xlsx"
out = r"c:\Users\PIXEL\OneDrive\Desktop\pos\tools\_odoo_bom2_analysis.txt"
print("exists", os.path.exists(path), "size", os.path.getsize(path))
wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
ws = wb.active
with open(out, "w", encoding="utf-8") as w:
    w.write("sheets: " + ", ".join(wb.sheetnames) + "\n")
    for i, row in enumerate(ws.iter_rows(max_row=8, values_only=True), start=1):
        w.write(f"R{i} ({len(row)} cols): {list(row)}\n")
    n = sum(1 for _ in ws.iter_rows(min_row=2, values_only=True))
    w.write(f"data_rows={n}\n")
with zipfile.ZipFile(path) as z:
    xml = z.read("xl/worksheets/sheet1.xml").decode("utf-8", errors="replace")
    m = re.search(r'ref="([^"]+)"', xml)
    if m:
        w = open(out, "a", encoding="utf-8")
        w.write(f"dimension={m.group(1)}\n")
        w.close()
print("done")
