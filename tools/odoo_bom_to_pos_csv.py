"""تحويل BOM من Odoo (Excel/CSV) إلى CSV استيراد POS."""
from __future__ import annotations

import csv
import re
import sys
from decimal import Decimal, InvalidOperation
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None  # type: ignore

csv.field_size_limit(sys.maxsize)

POS_BOM_HEADERS = [
    "parent_sku",
    "parent_name",
    "component_sku",
    "component_name",
    "qty_per_parent",
]

PARENT_KEYS = (
    "product",
    "parent",
    "parent_name",
    "parent product",
    "bom product",
    "finished product",
)
COMPONENT_KEYS = (
    "component",
    "component_name",
    "component product",
    "product (component)",
    "raw material",
    "ingredient",
    "product_id",
)
QTY_KEYS = (
    "qty_per_parent",
    "quantity",
    "product qty",
    "product_qty",
    "qty",
    "component quantity",
)
PARENT_SKU_KEYS = ("parent_sku", "parent reference", "parent internal reference")
COMPONENT_SKU_KEYS = (
    "component_sku",
    "component reference",
    "component internal reference",
)


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").strip())


def _norm_key(s: str) -> str:
    return re.sub(r"[\s_\-]+", " ", (s or "").strip().lower())


def _dec(raw: str) -> Decimal | None:
    try:
        v = Decimal((raw or "").strip().replace(",", "") or "0")
        return v if v > 0 else None
    except InvalidOperation:
        return None


def load_product_maps(path: Path) -> tuple[dict[str, tuple[str, str]], dict[str, str]]:
    """casefold(name) -> (sku, name_ar), sku -> name_ar"""
    by_name: dict[str, tuple[str, str]] = {}
    by_sku: dict[str, str] = {}
    if not path.is_file():
        return by_name, by_sku
    with path.open(encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            name = _norm(row.get("name_ar") or "")
            sku = _norm(row.get("sku") or "")
            if name and sku:
                by_name[name.casefold()] = (sku, name)
                by_sku[sku] = name
    return by_name, by_sku


def resolve_product(
    name: str, sku: str, by_name: dict[str, tuple[str, str]]
) -> tuple[str, str]:
    name = _norm(name)
    sku = _norm(sku)
    if sku:
        return sku, name
    if name:
        hit = by_name.get(name.casefold())
        if hit:
            return hit[0], hit[1]
    return sku, name


def _pick(row: dict[str, str], keys: tuple[str, ...]) -> str:
    norm_row = {_norm_key(k): (v or "") for k, v in row.items()}
    for key in keys:
        v = norm_row.get(_norm_key(key), "")
        if v.strip():
            return v.strip()
    return ""


def read_rows_from_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", errors="replace", newline="") as f:
        reader = csv.DictReader(f)
        return [dict(r) for r in reader]


def read_rows_from_xlsx(path: Path) -> list[dict[str, str]]:
    if openpyxl is None:
        raise RuntimeError("openpyxl غير مثبت.")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h or "").strip() for h in rows[0]]
    out: list[dict[str, str]] = []
    for raw in rows[1:]:
        row: dict[str, str] = {}
        for i, h in enumerate(headers):
            if not h:
                continue
            val = raw[i] if i < len(raw) else ""
            row[h] = "" if val is None else str(val).strip()
        if any(v.strip() for v in row.values()):
            out.append(row)
    return out


def is_odoo_hierarchical_export(rows: list[dict[str, str]]) -> bool:
    if not rows:
        return False
    keys = {_norm_key(k) for k in rows[0].keys()}
    return "product" in keys and any("bom lines" in k for k in keys)


def parse_odoo_hierarchical_rows(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    """تصدير Odoo: Product في سطر الوجبة + BoM Lines للمكوّنات."""
    flat: list[dict[str, str]] = []
    current_parent = ""
    for row in rows:
        parent_name = _pick(row, PARENT_KEYS)
        comp_name = _pick(row, ("bom lines", *COMPONENT_KEYS))
        if parent_name:
            current_parent = parent_name
            if comp_name:
                flat.append(
                    {
                        "Product": current_parent,
                        "Component": comp_name,
                        "Quantity": "1",
                    }
                )
        elif comp_name and current_parent:
            flat.append(
                {
                    "Product": current_parent,
                    "Component": comp_name,
                    "Quantity": "1",
                }
            )
    return flat


def convert_bom_rows(
    rows: list[dict[str, str]],
    products_csv: Path,
    *,
    default_qty: Decimal = Decimal("1"),
) -> tuple[list[dict[str, str]], list[str]]:
    by_name, _ = load_product_maps(products_csv)
    if is_odoo_hierarchical_export(rows):
        rows = parse_odoo_hierarchical_rows(rows)
    out: list[dict[str, str]] = []
    warnings: list[str] = []
    defaulted_qty = 0

    for i, row in enumerate(rows, start=2):
        parent_name = _pick(row, PARENT_KEYS)
        parent_sku = _pick(row, PARENT_SKU_KEYS)
        comp_name = _pick(row, COMPONENT_KEYS)
        comp_sku = _pick(row, COMPONENT_SKU_KEYS)
        qty_raw = _pick(row, QTY_KEYS)

        if not comp_name and not comp_sku:
            if parent_name:
                warnings.append(
                    f"سطر {i}: وجبة «{parent_name}» بدون مكوّn — الملف يحتوي رؤوس BOM فقط."
                )
            continue

        qty = _dec(qty_raw)
        if qty is None:
            qty = default_qty
            defaulted_qty += 1

        parent_sku, parent_name = resolve_product(parent_name, parent_sku, by_name)
        comp_sku, comp_name = resolve_product(comp_name, comp_sku, by_name)

        if not parent_name and not parent_sku:
            warnings.append(f"سطر {i}: الوجبة (الأب) غير محددة.")
            continue
        if not comp_name and not comp_sku:
            warnings.append(f"سطر {i}: المكوّn غير محدد.")
            continue
        if not parent_sku:
            warnings.append(
                f"سطر {i}: لم يُعثر على SKU للوجبة «{parent_name}» في products_pos_import.csv."
            )
        if not comp_sku:
            warnings.append(
                f"سطر {i}: لم يُعثر على SKU للمكوّn «{comp_name}» في products_pos_import.csv."
            )

        out.append(
            {
                "parent_sku": parent_sku,
                "parent_name": parent_name,
                "component_sku": comp_sku,
                "component_name": comp_name,
                "qty_per_parent": str(qty.quantize(Decimal("0.001"))),
            }
        )

    if defaulted_qty:
        warnings.insert(
            0,
            f"⚠️ {defaulted_qty} سطراً بدون كمية في Odoo — وُضعت qty=1 (راجع الكميات بالجرام/اللتر).",
        )
    return out, warnings


def write_bom_csv(path: Path, rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=POS_BOM_HEADERS)
        w.writeheader()
        for row in rows:
            w.writerow(row)


def count_parents_in_hierarchical(rows: list[dict[str, str]]) -> int:
    if not is_odoo_hierarchical_export(rows):
        return 0
    return sum(1 for r in rows if _pick(r, PARENT_KEYS))


def convert_file(src: Path, dst: Path, products_csv: Path) -> tuple[int, list[str], int]:
    if src.suffix.lower() in {".xlsx", ".xlsm", ".xltx"}:
        source_rows = read_rows_from_xlsx(src)
    else:
        source_rows = read_rows_from_csv(src)
    parent_count = count_parents_in_hierarchical(source_rows)
    out, warnings = convert_bom_rows(source_rows, products_csv)
    write_bom_csv(dst, out)
    return len(out), warnings, parent_count


def main() -> None:
    base = Path(r"c:\Users\PIXEL\OneDrive\Desktop\سيف\New folder")
    src = base / "Bill of Material (mrp.bom) (1).xlsx"
    products_csv = base / "products_pos_import.csv"
    dst = base / "bom_pos_import.csv"
    summary = base / "bom_pos_import_summary.txt"

    if len(sys.argv) >= 2:
        src = Path(sys.argv[1])
    if len(sys.argv) >= 3:
        dst = Path(sys.argv[2])

    count, warnings, parent_count = convert_file(src, dst, products_csv)

    lines = [
        f"source={src}",
        f"products_map={products_csv}",
        f"output={dst}",
        f"bom_lines={count}",
        f"bom_parents={parent_count}",
        "",
    ]
    if count == 0:
        lines.append("⚠️ لم يُستورد أي سطر BOM.")
    else:
        lines.append("✅ جاهز للرفع من /catalog/import-export → وصفات التركيب (BOM)")
    if warnings:
        lines.append("")
        lines.append(f"warnings ({len(warnings)}):")
        lines.extend(warnings[:80])
        if len(warnings) > 80:
            lines.append(f"... و{len(warnings)-80} تحذيراً آخر")

    summary.write_text("\n".join(lines), encoding="utf-8")
    print(f"bom_lines={count} parents={parent_count}")


if __name__ == "__main__":
    main()
