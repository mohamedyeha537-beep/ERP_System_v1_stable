"""تصدير التقارير إلى CSV و Excel (.xlsx).

CSV: UTF-8-BOM لعرض العربية في Excel على Windows.
XLSX: openpyxl — ترميز Unicode أصلي دون تحويل إلى رموز.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from typing import Iterable, Sequence


_CSV_BOM = "\ufeff"


def _fmt(value) -> str:
    """تحويل أي قيمة إلى نص مناسب لـ CSV."""
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return f"{value:.3f}"
    if isinstance(value, datetime):
        from app.datetime_local import format_local_dt

        return format_local_dt(value, "%Y-%m-%d %H:%M")
    return str(value)


def make_csv(
    headers: Sequence[str], rows: Iterable[Sequence]
) -> str:
    """يبني نص CSV مع BOM.
    headers: ['اسم', 'كمية', ...]
    rows: قائمة من tuples/lists.
    """
    buf = io.StringIO()
    buf.write(_CSV_BOM)
    writer = csv.writer(buf, lineterminator="\n", quoting=csv.QUOTE_MINIMAL)
    writer.writerow(headers)
    for r in rows:
        writer.writerow([_fmt(c) for c in r])
    return buf.getvalue()


def csv_response(
    filename: str,
    headers: Sequence[str],
    rows: Iterable[Sequence],
):
    """يُعيد Response جاهزاً لـ FastAPI كمرفق تنزيل."""
    from fastapi.responses import Response

    body = make_csv(headers, rows).encode("utf-8")
    return Response(
        content=body,
        media_type="text/csv; charset=utf-8",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}.csv"',
            "Content-Length": str(len(body)),
        },
    )


def _cell_value(value) -> str | int | float:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, (int, float)):
        return value
    text = str(value)
    if re.fullmatch(r"-?\d+\.\d+", text):
        try:
            return float(text)
        except ValueError:
            pass
    return text


@dataclass
class SheetSpec:
    name: str
    headers: Sequence[str]
    rows: Iterable[Sequence]


def make_xlsx(sheets: Sequence[SheetSpec]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    header_font = Font(bold=True)

    for spec in sheets:
        title = (spec.name or "تقرير")[:31]
        ws = wb.create_sheet(title=title)
        ws.sheet_view.rightToLeft = True

        for col_idx, header in enumerate(spec.headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)

        for row_idx, row in enumerate(spec.rows, start=2):
            for col_idx, raw in enumerate(row, start=1):
                val = _cell_value(raw)
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                if isinstance(val, (int, float)):
                    cell.number_format = "#,##0.000"
                    cell.alignment = Alignment(horizontal="left")
                else:
                    cell.alignment = Alignment(horizontal="right", wrap_text=True)

        for col_idx in range(1, len(spec.headers) + 1):
            max_len = len(str(spec.headers[col_idx - 1]))
            for row in ws.iter_rows(
                min_row=2, max_row=ws.max_row, min_col=col_idx, max_col=col_idx
            ):
                for cell in row:
                    if cell.value is not None:
                        max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 48)

    if not wb.sheetnames:
        ws = wb.create_sheet("تقرير")
        ws.sheet_view.rightToLeft = True

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xlsx_response(filename: str, sheets: Sequence[SheetSpec]):
    from fastapi.responses import Response

    body = make_xlsx(sheets)
    safe = re.sub(r"[^\w\-]+", "-", filename).strip("-") or "report"
    return Response(
        content=body,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{safe}.xlsx"',
            "Content-Length": str(len(body)),
        },
    )
